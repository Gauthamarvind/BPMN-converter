"""
Unified Ingestion Engine.
Parses transcripts (.txt, .md, .vtt, .srt), documents (.docx, .pdf),
and spreadsheets (.xlsx, .csv) into normalized text and structured tables.
Raises IngestionError on corrupt files, empty text, or unreadable PDFs.
"""

from __future__ import annotations
import io
import re
import csv
from pathlib import Path
from dataclasses import dataclass, field
from typing import List, Dict, Any, Optional, Tuple, Union


class IngestionError(Exception):
    """Raised when document ingestion fails due to corrupt files, unreadable formats, or empty content."""
    pass


@dataclass
class IngestedDocument:
    filename: str
    extension: str
    normalized_text: str
    tables: List[List[Dict[str, Any]]] = field(default_factory=list)
    metadata: Dict[str, Any] = field(default_factory=dict)


def clean_vtt_srt(content: str) -> str:
    """Strips timestamps, sequence numbers, and formatting from VTT / SRT files."""
    lines = content.splitlines()
    cleaned_lines = []

    # Regex to match timestamps like "00:01:20.000 --> 00:01:23.000" or "00:01:20,000 --> 00:01:23,000"
    timestamp_pattern = re.compile(r"\d{1,2}:\d{2}:\d{2}[.,]\d{3}\s*-->\s*\d{1,2}:\d{2}:\d{2}[.,]\d{3}")

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("WEBVTT") or stripped.startswith("NOTE"):
            continue
        if stripped.isdigit():  # Sequence counter in SRT
            continue
        if timestamp_pattern.search(stripped):
            continue

        # Strip XML-like tags e.g. <v Speaker> or <c.color>
        speaker_match = re.search(r"<v\s+([^>]+)>", stripped)
        if speaker_match:
            speaker = speaker_match.group(1)
            text = re.sub(r"<[^>]+>", "", stripped)
            stripped = f"{speaker}: {text.strip()}"
        else:
            stripped = re.sub(r"<[^>]+>", "", stripped)

        cleaned_lines.append(stripped)

    return "\n".join(cleaned_lines)


def parse_csv_content(content: str) -> Tuple[str, List[Dict[str, Any]]]:
    """Parses CSV with auto-detection of process-relevant columns."""
    f = io.StringIO(content)
    try:
        sample = content[:2048]
        delimiter = ","
        if sample:
            sniffer = csv.Sniffer()
            dialect = sniffer.sniff(sample)
            delimiter = dialect.delimiter
    except Exception:
        delimiter = ","

    f.seek(0)
    reader = csv.reader(f, delimiter=delimiter)
    rows = list(reader)
    if not rows:
        return "", []

    header = [c.strip() for c in rows[0]]
    column_mapping = detect_column_roles(header)

    structured_rows = []
    text_lines = [f"=== CSV Process Table ===", f"Columns: {', '.join(header)}"]

    for row_idx, r in enumerate(rows[1:], start=1):
        if not any(cell.strip() for cell in r):
            continue
        row_dict = {}
        text_parts = []
        for col_idx, cell in enumerate(r):
            val = cell.strip()
            standard_role = column_mapping.get(col_idx, header[col_idx] if col_idx < len(header) else f"col_{col_idx}")
            row_dict[standard_role] = val
            if val:
                text_parts.append(f"{standard_role}: {val}")

        structured_rows.append(row_dict)
        text_lines.append(f"Row {row_idx} -> " + " | ".join(text_parts))

    return "\n".join(text_lines), structured_rows


def detect_column_roles(headers: List[str]) -> Dict[int, str]:
    """Fuzzy column mapper for business process spreadsheets."""
    mapping = {}
    for idx, h in enumerate(headers):
        hl = h.lower().strip()
        if any(k in hl for k in ["next", "successor", "then", "following", "leads to"]):
            mapping[idx] = "next_step"
        elif any(k in hl for k in ["id", "no.", "number", "#"]):
            mapping[idx] = "step_id"
        elif any(k in hl for k in ["role", "actor", "lane", "owner", "performer", "who", "department", "user"]):
            mapping[idx] = "actor_role"
        elif any(k in hl for k in ["condition", "decision", "rule", "criteria", "branch", "if"]):
            mapping[idx] = "condition"
        elif any(k in hl for k in ["system", "application", "tool", "software", "service", "platform"]):
            mapping[idx] = "system"
        elif any(k in hl for k in ["input", "data in", "trigger"]):
            mapping[idx] = "input_data"
        elif any(k in hl for k in ["output", "data out", "result", "artifact"]):
            mapping[idx] = "output_data"
        elif any(k in hl for k in ["time", "duration", "sla", "delay"]):
            mapping[idx] = "duration"
        elif any(k in hl for k in ["step", "activity", "task", "action", "description", "event", "process"]):
            mapping[idx] = "step_name"
        else:
            mapping[idx] = hl.replace(" ", "_")
    return mapping


def parse_xlsx_bytes(data: bytes) -> Tuple[str, List[List[Dict[str, Any]]]]:
    """Parses Excel workbook sheets and extracts structured process data."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as ex:
        raise IngestionError(f"Corrupt or invalid Excel workbook: {str(ex)}") from ex

    all_text = []
    all_tables = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        non_empty_rows = [r for r in rows if any(c is not None and str(c).strip() for c in r)]
        if not non_empty_rows:
            continue

        header = [str(c).strip() if c is not None else f"Col_{i}" for i, c in enumerate(non_empty_rows[0])]
        column_mapping = detect_column_roles(header)

        sheet_text = [f"=== Sheet: {sheet_name} ===", f"Columns: {', '.join(header)}"]
        table_rows = []

        for row_idx, r in enumerate(non_empty_rows[1:], start=1):
            row_dict = {}
            parts = []
            for col_idx, cell in enumerate(r):
                val = str(cell).strip() if cell is not None else ""
                col_name = header[col_idx] if col_idx < len(header) else f"Col_{col_idx}"
                std_name = column_mapping.get(col_idx, col_name)
                row_dict[std_name] = val
                if val:
                    parts.append(f"{std_name}: {val}")
            table_rows.append(row_dict)
            sheet_text.append(f"Row {row_idx}: " + " | ".join(parts))

        all_text.append("\n".join(sheet_text))
        all_tables.append(table_rows)

    return "\n\n".join(all_text), all_tables


def parse_docx_bytes(data: bytes) -> Tuple[str, List[List[Dict[str, Any]]]]:
    """Parses Word .docx document paragraphs, headings, bullet lists, and tables."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(data))
    except Exception as ex:
        raise IngestionError(f"Corrupt or invalid Word document: {str(ex)}") from ex

    lines = []
    # Extract paragraphs
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            # Check style name for headings or lists
            style_name = p.style.name.lower() if p.style else ""
            if "heading 1" in style_name:
                lines.append(f"\n# {text}")
            elif "heading 2" in style_name:
                lines.append(f"\n## {text}")
            elif "heading" in style_name:
                lines.append(f"\n### {text}")
            elif "list" in style_name or "bullet" in style_name:
                lines.append(f"- {text}")
            else:
                lines.append(text)

    # Extract tables in docx
    tables = []
    for t_idx, table in enumerate(doc.tables, start=1):
        t_rows = []
        for row in table.rows:
            t_rows.append([cell.text.strip() for cell in row.cells])
        if t_rows:
            header = t_rows[0]
            col_map = detect_column_roles(header)
            table_data = []
            table_lines = [f"\n--- Document Table {t_idx} ---", f"Columns: {', '.join(header)}"]
            for row in t_rows[1:]:
                row_dict = {}
                parts = []
                for idx, cell in enumerate(row):
                    std = col_map.get(idx, f"col_{idx}")
                    row_dict[std] = cell
                    if cell:
                        parts.append(f"{std}: {cell}")
                table_data.append(row_dict)
                table_lines.append(" | ".join(parts))
            tables.append(table_data)
            lines.append("\n".join(table_lines))

    return "\n".join(lines), tables


def parse_pdf_bytes(data: bytes) -> str:
    """Extracts text page-by-page from PDF documents."""
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(data))
        pages_text = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages_text.append(f"--- Page {i + 1} ---\n{text.strip()}")
        if not pages_text:
            raise IngestionError("PDF document contains no extractable text.")
        return "\n\n".join(pages_text)
    except IngestionError:
        raise
    except Exception as ex:
        raise IngestionError(f"Corrupt or invalid PDF document: {str(ex)}") from ex


def ingest_file(
    content: Union[str, bytes],
    filename: str
) -> IngestedDocument:
    """
    Main ingestion dispatcher. Normalizes any input format into structured process text.
    Raises IngestionError if document is corrupt or has empty content.
    """
    ext = Path(filename).suffix.lower()
    meta = {"source_filename": filename, "format": ext}
    tables: List[List[Dict[str, Any]]] = []

    # Text formats
    if ext in (".txt", ".md", ".markdown"):
        raw_text = content.decode("utf-8", errors="ignore") if isinstance(content, bytes) else str(content)
        normalized = raw_text.strip()
    elif ext in (".vtt", ".srt"):
        raw_text = content.decode("utf-8", errors="ignore") if isinstance(content, bytes) else str(content)
        normalized = clean_vtt_srt(raw_text)
    elif ext == ".csv":
        raw_text = content.decode("utf-8", errors="ignore") if isinstance(content, bytes) else str(content)
        normalized, table = parse_csv_content(raw_text)
        if table:
            tables.append(table)
    elif ext == ".docx":
        b_data = content if isinstance(content, bytes) else content.encode("utf-8")
        normalized, tables = parse_docx_bytes(b_data)
    elif ext == ".pdf":
        b_data = content if isinstance(content, bytes) else content.encode("utf-8")
        normalized = parse_pdf_bytes(b_data)
    elif ext in (".xlsx", ".xls"):
        b_data = content if isinstance(content, bytes) else content.encode("utf-8")
        normalized, tables = parse_xlsx_bytes(b_data)
    else:
        # Fallback raw text decoding
        raw_text = content.decode("utf-8", errors="ignore") if isinstance(content, bytes) else str(content)
        normalized = raw_text.strip()

    if not normalized or not normalized.strip():
        raise IngestionError(f"Ingested document '{filename}' is empty or contains no readable text.")

    return IngestedDocument(
        filename=filename,
        extension=ext,
        normalized_text=normalized,
        tables=tables,
        metadata=meta
    )
