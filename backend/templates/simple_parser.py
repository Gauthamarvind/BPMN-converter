"""
Simple Document & Table Template Parser (v2).
Implements deterministic mapping from simple table/spreadsheet steps to ProcessIR.
Supports: .xlsx, .docx, .csv, and .json formats.
Enforces row-level validations, signature-based template detection, and sourceRefs.
"""

from __future__ import annotations
import csv
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Any, Set

import openpyxl
import docx
from pydantic import BaseModel, Field

from backend.ir.models import (
    ProcessIR,
    FlowNode,
    SequenceFlow,
    Pool,
    Lane,
    SourceRef,
    OpenQuestion,
    DataObject
)
from backend.pipeline.validator import sanitize_ncname


TEMPLATE_KIND_SIMPLE = "simple"
TEMPLATE_KIND_LEGACY = "legacy"
TEMPLATE_KIND_NONE = "none"

SIMPLE_HEADERS_CORE = ["step id", "step", "responsible", "type", "if yes", "if no", "parallel group", "next step"]
LEGACY_HEADERS_CORE = ["step #", "activity", "role", "task type", "decision?", "next steps"]


@dataclass
class RowIssue:
    row: int
    column: str
    message: str
    severity: str = "ERROR"  # "ERROR", "WARNING", "INFO"
    fix: Optional[str] = None


def _source_ref(doc_id: str, sheet_name: str, row_number: int, text: str) -> SourceRef:
    """Builds a SourceRef in the IR's field layout: '<file>!<sheet>:row <n>' plus the row text."""
    location = f"{doc_id}!{sheet_name}:row {row_number}" if doc_id else f"{sheet_name}:row {row_number}"
    return SourceRef(sourceLocation=location, textSnippet=(text or "")[:200])


class RowValidationError(ValueError):
    """Raised when one or more row validation errors prevent conversion."""
    def __init__(self, issues: List[RowIssue]):
        self.issues = issues
        self.errors = [i for i in issues if i.severity == "ERROR"]
        self.warnings = [i for i in issues if i.severity == "WARNING"]
        super().__init__(f"Template validation failed with {len(self.errors)} error(s).")


class StepRow(BaseModel):
    step_id: str
    step: str
    responsible: str = "Unassigned"
    type: str = "Task"  # "Task", "Decision", "End"
    if_yes: Optional[str] = None
    if_no: Optional[str] = None
    parallel_group: Optional[str] = None
    next_step: Optional[str] = None
    description: Optional[str] = None
    system: Optional[str] = None
    input_data: Optional[str] = None
    output_data: Optional[str] = None
    duration: Optional[str] = None
    row_number: int = 1
    sheet_name: str = "Process"


def normalize_str(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


# Column detection: whole-word / phrase matching, evaluated most-specific first, so a column
# called "Notes" no longer lands in "If No" and "Input" never matches "Output".
_HEADER_RULES: List[Tuple[str, Tuple[str, ...]]] = [
    ("step_id",        ("step id", "step no", "step number", "step #", "id", "no.", "#")),
    ("if_yes",         ("if yes", "if yes step", "yes", "yes step", "on yes", "true")),
    ("if_no",          ("if no", "if no step", "no", "no step", "on no", "false")),
    ("parallel_group", ("parallel group", "parallel", "group", "concurrent")),
    ("next_step",      ("next step", "next", "next steps", "successor", "goto")),
    ("responsible",    ("responsible", "role", "actor", "owner", "performer", "lane", "department", "who")),
    ("type",           ("type", "step type", "task type", "kind")),
    ("description",    ("description", "details", "instructions", "notes")),
    ("system",         ("system", "system/tool", "tool", "application", "app")),
    ("input_data",     ("input", "inputs", "input data")),
    ("output_data",    ("output", "outputs", "output data")),
    ("duration",       ("duration", "sla", "time", "estimate")),
    ("step",           ("step", "activity", "activity name", "task", "action", "name")),
]


def _normalize_header(name: Any) -> str:
    txt = normalize_str(name).lower().replace("→", "->").replace("–", "-").replace("—", "-")
    txt = re.sub(r"\s*->\s*", " -> ", txt)
    txt = re.sub(r"[\s_]+", " ", txt).strip()
    return txt


def map_header_columns(names: List[Any], offset: int = 0) -> Dict[str, int]:
    """
    Maps a header row to canonical StepRow field names.

    A header matches a rule when, after normalisation, it *equals* one of the rule's
    phrases or starts with it followed by a separator (so "If Yes -> Step" and "If Yes
    (step id)" match "if yes", but "Notes" does not match "no"). Rules are ordered
    most-specific first and each canonical field is assigned once.
    """
    col_map: Dict[str, int] = {}
    for idx, raw in enumerate(names):
        h = _normalize_header(raw)
        if not h:
            continue
        for field, phrases in _HEADER_RULES:
            if field in col_map:
                continue
            hit = False
            for ph in phrases:
                if h == ph or h.startswith(ph + " ") or h.startswith(ph + "-") or h.startswith(ph + "("):
                    hit = True
                    break
            if hit:
                col_map[field] = idx + offset
                break
    return col_map


def _looks_like_legacy_header(cells: List[str]) -> bool:
    """
    A step table with a step column, an actor/role column and an activity or next-step column
    (e.g. "Step ID | Actor / Role | Activity Name | Next Step") is a legacy step list.
    """
    norm = [re.sub(r"[^a-z0-9]+", " ", c).strip() for c in cells]
    has_step = any(c in ("step", "step id", "id", "no", "index") or c.startswith("step ") for c in norm)
    has_actor = any(any(tok in c.split() for tok in ("actor", "role", "lane", "department", "performer", "owner")) for c in norm)
    has_activity = any(any(tok in c.split() for tok in ("activity", "task", "action")) or c.startswith("next") for c in norm)
    return has_step and has_actor and has_activity


def detect_template_kind(raw_content: bytes, filename: str) -> str:
    """
    Detects template signature based on content rather than simple file extension.
    Returns 'simple', 'legacy', or 'none'.
    """
    if not raw_content:
        return TEMPLATE_KIND_NONE

    ext = Path(filename).suffix.lower()

    # 1. Excel Workbook (.xlsx)
    if ext in (".xlsx", ".xls") or (ext not in (".docx", ".pdf", ".csv", ".json") and raw_content.startswith(b"PK") and b"xl/" in raw_content[:4096]):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw_content), read_only=True, data_only=True)
            # Check hidden _meta sheet
            if "_meta" in wb.sheetnames:
                ws_meta = wb["_meta"]
                for row in ws_meta.iter_rows(values_only=True):
                    if row and len(row) >= 2:
                        k, v = normalize_str(row[0]).lower(), normalize_str(row[1])
                        if k == "template_version" and v == "2":
                            wb.close()
                            return TEMPLATE_KIND_SIMPLE

            # Check sheet header row (Process sheet or first sheet)
            sheet_to_check = "Process" if "Process" in wb.sheetnames else wb.sheetnames[0]
            ws = wb[sheet_to_check]
            for row in ws.iter_rows(values_only=True):
                if not row:
                    continue
                row_str_cells = [normalize_str(c).lower() for c in row if c is not None]
                joined = " ".join(row_str_cells)
                if not joined:
                    continue
                # Simple parser header detection
                if "step id" in joined and "responsible" in joined:
                    wb.close()
                    return TEMPLATE_KIND_SIMPLE
                if "step" in row_str_cells and "responsible" in row_str_cells and "type" in row_str_cells:
                    wb.close()
                    return TEMPLATE_KIND_SIMPLE
                # Legacy parser header detection
                if "step #" in joined or ("activity" in row_str_cells and "role" in row_str_cells and "task type" in row_str_cells):
                    wb.close()
                    return TEMPLATE_KIND_LEGACY
                if _looks_like_legacy_header(row_str_cells):
                    wb.close()
                    return TEMPLATE_KIND_LEGACY
            wb.close()
        except Exception:
            pass

    # 2. Word Document (.docx)
    if ext == ".docx":
        try:
            doc = docx.Document(io.BytesIO(raw_content))
            for table in doc.tables:
                if not table.rows:
                    continue
                hdr_cells = [normalize_str(c.text).lower() for c in table.rows[0].cells]
                joined = " ".join(hdr_cells)
                if "step id" in joined and "responsible" in joined:
                    return TEMPLATE_KIND_SIMPLE
                if "step" in hdr_cells and "responsible" in hdr_cells and "type" in hdr_cells:
                    return TEMPLATE_KIND_SIMPLE
                if "step #" in joined or ("activity" in hdr_cells and "role" in hdr_cells and "task type" in hdr_cells):
                    return TEMPLATE_KIND_LEGACY
            # If docx has no matching table header, it is an SOP document!
            return TEMPLATE_KIND_NONE
        except Exception:
            pass

    # 3. CSV (.csv)
    if ext == ".csv":
        try:
            text = raw_content.decode("utf-8-sig", errors="ignore")
            reader = csv.reader(io.StringIO(text))
            for row in reader:
                if not row:
                    continue
                row_str_cells = [normalize_str(c).lower() for c in row]
                joined = " ".join(row_str_cells)
                if "step id" in joined and "responsible" in joined:
                    return TEMPLATE_KIND_SIMPLE
                if "step" in row_str_cells and "responsible" in row_str_cells:
                    return TEMPLATE_KIND_SIMPLE
                if "step" in row_str_cells and ("task" in row_str_cells or "actor" in row_str_cells or "next steps" in row_str_cells or "type" in row_str_cells):
                    return TEMPLATE_KIND_LEGACY
                if "step #" in joined or ("activity" in row_str_cells and "role" in row_str_cells) or "task type" in row_str_cells:
                    return TEMPLATE_KIND_LEGACY
                if _looks_like_legacy_header(row_str_cells):
                    return TEMPLATE_KIND_LEGACY
        except Exception:
            pass

    # 4. JSON (.json)
    if ext == ".json" or raw_content.strip().startswith(b"{"):
        try:
            data = json.loads(raw_content.decode("utf-8", errors="ignore"))
            if isinstance(data, dict):
                if data.get("template_version") == 2 or "steps" in data:
                    if "steps" in data and isinstance(data["steps"], list) and len(data["steps"]) > 0:
                        first = data["steps"][0]
                        if isinstance(first, dict) and ("step" in first or "step_id" in first or "responsible" in first):
                            return TEMPLATE_KIND_SIMPLE
                if "parsing_rules" in data or "columns" in data:
                    return TEMPLATE_KIND_LEGACY
        except Exception:
            pass

    return TEMPLATE_KIND_NONE


class SimpleTemplateParser:
    """
    Deterministic parser for the standard Process2BPMN v2 template.
    Maps steps directly into a valid ProcessIR model.
    """

    def __init__(self):
        pass

    def parse_bytes(self, raw_content: bytes, filename: str) -> ProcessIR:
        """Parses byte payload into a ProcessIR or raises RowValidationError / ValueError."""
        ext = Path(filename).suffix.lower()
        process_name = Path(filename).stem.replace("_", " ").title()

        if ext in (".xlsx", ".xls") or (ext not in (".docx", ".pdf", ".csv", ".json") and raw_content.startswith(b"PK") and b"xl/" in raw_content[:4096]):
            return self._parse_xlsx(raw_content, process_name, filename)
        elif ext == ".docx" or (ext not in (".xlsx", ".pdf", ".csv", ".json") and raw_content.startswith(b"PK") and b"word/" in raw_content[:4096]):
            return self._parse_docx(raw_content, process_name, filename)
        elif ext == ".csv":
            return self._parse_csv(raw_content, process_name, filename)
        elif ext == ".json":
            return self._parse_json(raw_content, process_name, filename)
        else:
            raise ValueError(f"Unsupported format '{ext}' for SimpleTemplateParser.")

    def _parse_xlsx(self, raw_bytes: bytes, default_name: str, filename: str) -> ProcessIR:
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
        sheet_name = "Process" if "Process" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]

        process_name = default_name
        # Look for process name in cell B1 or A1
        cell_a1 = normalize_str(ws["A1"].value)
        cell_b1 = normalize_str(ws["B1"].value)
        if "process name" in cell_a1.lower() and cell_b1:
            process_name = cell_b1
        elif cell_a1 and not cell_a1.lower().startswith("step"):
            process_name = cell_a1

        # Locate header row
        header_row_idx = -1
        col_map: Dict[str, int] = {}

        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if not row:
                continue
            row_lower = [normalize_str(c).lower() for c in row]
            if any("step id" in c or c == "step" for c in row_lower) and any("responsible" in c or "type" in c for c in row_lower):
                header_row_idx = r_idx
                col_map = map_header_columns(list(row), offset=1)
                break

        if header_row_idx == -1:
            raise ValueError(f"Could not locate table header in sheet '{sheet_name}'.")

        steps: List[StepRow] = []
        for r_idx in range(header_row_idx + 1, ws.max_row + 1):
            step_id_val = ws.cell(row=r_idx, column=col_map.get("step_id", 1)).value
            step_text_val = ws.cell(row=r_idx, column=col_map.get("step", 2)).value
            if step_id_val is None and step_text_val is None:
                continue

            step_id = normalize_str(step_id_val)
            step_text = normalize_str(step_text_val)
            if not step_id and not step_text:
                continue

            resp = normalize_str(ws.cell(row=r_idx, column=col_map.get("responsible", 3)).value) or ""
            t_val = normalize_str(ws.cell(row=r_idx, column=col_map.get("type", 4)).value) or "Task"
            if_yes = normalize_str(ws.cell(row=r_idx, column=col_map.get("if_yes", 5)).value) or None
            if_no = normalize_str(ws.cell(row=r_idx, column=col_map.get("if_no", 6)).value) or None
            p_group = normalize_str(ws.cell(row=r_idx, column=col_map.get("parallel_group", 7)).value) or None
            next_step = normalize_str(ws.cell(row=r_idx, column=col_map.get("next_step", 8)).value) or None

            desc = normalize_str(ws.cell(row=r_idx, column=col_map.get("description", 9)).value) if "description" in col_map else None
            sys = normalize_str(ws.cell(row=r_idx, column=col_map.get("system", 10)).value) if "system" in col_map else None
            inp = normalize_str(ws.cell(row=r_idx, column=col_map.get("input_data", 11)).value) if "input_data" in col_map else None
            out = normalize_str(ws.cell(row=r_idx, column=col_map.get("output_data", 12)).value) if "output_data" in col_map else None
            dur = normalize_str(ws.cell(row=r_idx, column=col_map.get("duration", 13)).value) if "duration" in col_map else None

            steps.append(StepRow(
                step_id=step_id or str(len(steps) + 1),
                step=step_text or f"Step {step_id}",
                responsible=resp,
                type=t_val,
                if_yes=if_yes,
                if_no=if_no,
                parallel_group=p_group,
                next_step=next_step,
                description=desc,
                system=sys,
                input_data=inp,
                output_data=out,
                duration=dur,
                row_number=r_idx,
                sheet_name=sheet_name
            ))

        return self.build_process_ir(steps, process_name=process_name, doc_id=filename)

    def _parse_docx(self, raw_bytes: bytes, default_name: str, filename: str) -> ProcessIR:
        doc = docx.Document(io.BytesIO(raw_bytes))
        process_name = default_name
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                if "process:" in txt.lower() or "process name:" in txt.lower():
                    parts = txt.split(":", 1)
                    if len(parts) > 1 and parts[1].strip():
                        process_name = parts[1].strip()
                        break
                elif not process_name or process_name == default_name:
                    process_name = txt
                    break

        table = None
        for t in doc.tables:
            if t.rows:
                hdr = [normalize_str(c.text).lower() for c in t.rows[0].cells]
                if any("step" in c for c in hdr) and any("responsible" in c or "type" in c for c in hdr):
                    table = t
                    break

        if not table:
            raise ValueError("No valid process table found in Word document.")

        col_map: Dict[str, int] = map_header_columns([c.text for c in table.rows[0].cells])

        steps: List[StepRow] = []
        for r_idx, row in enumerate(table.rows[1:], start=2):
            cells = [normalize_str(c.text) for c in row.cells]
            if not any(cells):
                continue
            step_id = cells[col_map.get("step_id", 0)] if col_map.get("step_id", 0) < len(cells) else str(r_idx - 1)
            step_text = cells[col_map.get("step", 1)] if col_map.get("step", 1) < len(cells) else ""
            if not step_text and not step_id:
                continue

            resp = cells[col_map["responsible"]] if "responsible" in col_map and col_map["responsible"] < len(cells) else "Unassigned"
            t_val = cells[col_map["type"]] if "type" in col_map and col_map["type"] < len(cells) else "Task"
            if_yes = cells[col_map["if_yes"]] if "if_yes" in col_map and col_map["if_yes"] < len(cells) else None
            if_no = cells[col_map["if_no"]] if "if_no" in col_map and col_map["if_no"] < len(cells) else None
            p_group = cells[col_map["parallel_group"]] if "parallel_group" in col_map and col_map["parallel_group"] < len(cells) else None
            next_step = cells[col_map["next_step"]] if "next_step" in col_map and col_map["next_step"] < len(cells) else None

            steps.append(StepRow(
                step_id=step_id or str(len(steps) + 1),
                step=step_text or f"Step {step_id}",
                responsible=resp or "Unassigned",
                type=t_val or "Task",
                if_yes=if_yes or None,
                if_no=if_no or None,
                parallel_group=p_group or None,
                next_step=next_step or None,
                row_number=r_idx,
                sheet_name="Document"
            ))

        return self.build_process_ir(steps, process_name=process_name, doc_id=filename)

    def _parse_csv(self, raw_bytes: bytes, default_name: str, filename: str) -> ProcessIR:
        text = raw_bytes.decode("utf-8-sig", errors="ignore")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        if not rows:
            raise ValueError("CSV file is empty.")

        header_idx = -1
        col_map: Dict[str, int] = {}
        for r_idx, row in enumerate(rows, start=1):
            row_lower = [normalize_str(c).lower().replace("→", "->").replace("–", "-") for c in row]
            if any("step" in c for c in row_lower) and any("responsible" in c or "type" in c for c in row_lower):
                header_idx = r_idx - 1
                col_map = map_header_columns(list(row))
                break

        if header_idx == -1:
            raise ValueError("Could not find table header row in CSV.")

        steps: List[StepRow] = []
        for r_idx, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
            if not row or not any(row):
                continue
            cells = [normalize_str(c) for c in row]
            step_id = cells[col_map.get("step_id", 0)] if col_map.get("step_id", 0) < len(cells) else str(len(steps) + 1)
            step_text = cells[col_map.get("step", 1)] if col_map.get("step", 1) < len(cells) else ""
            if not step_id and not step_text:
                continue

            resp = cells[col_map["responsible"]] if "responsible" in col_map and col_map["responsible"] < len(cells) else "Unassigned"
            t_val = cells[col_map["type"]] if "type" in col_map and col_map["type"] < len(cells) else "Task"
            if_yes = cells[col_map["if_yes"]] if "if_yes" in col_map and col_map["if_yes"] < len(cells) else None
            if_no = cells[col_map["if_no"]] if "if_no" in col_map and col_map["if_no"] < len(cells) else None
            p_group = cells[col_map["parallel_group"]] if "parallel_group" in col_map and col_map["parallel_group"] < len(cells) else None
            next_step = cells[col_map["next_step"]] if "next_step" in col_map and col_map["next_step"] < len(cells) else None

            steps.append(StepRow(
                step_id=step_id or str(len(steps) + 1),
                step=step_text or f"Step {step_id}",
                responsible=resp or "Unassigned",
                type=t_val or "Task",
                if_yes=if_yes or None,
                if_no=if_no or None,
                parallel_group=p_group or None,
                next_step=next_step or None,
                row_number=r_idx,
                sheet_name="CSV"
            ))

        return self.build_process_ir(steps, process_name=default_name, doc_id=filename)

    def _parse_json(self, raw_bytes: bytes, default_name: str, filename: str) -> ProcessIR:
        data = json.loads(raw_bytes.decode("utf-8", errors="ignore"))
        process_name = data.get("process_name") or default_name
        steps_list = data.get("steps", [])
        if not isinstance(steps_list, list) or not steps_list:
            raise ValueError("JSON template payload must contain a non-empty 'steps' list.")

        steps: List[StepRow] = []
        for idx, item in enumerate(steps_list, start=1):
            if not isinstance(item, dict):
                continue
            step_id = normalize_str(item.get("step_id") or item.get("id") or str(idx))
            step_text = normalize_str(item.get("step") or item.get("name") or item.get("activity") or f"Step {step_id}")
            resp = normalize_str(item.get("responsible") or item.get("role") or "Unassigned")
            t_val = normalize_str(item.get("type") or item.get("task_type") or "Task")
            if_yes = normalize_str(item.get("if_yes") or item.get("yes")) or None
            if_no = normalize_str(item.get("if_no") or item.get("no")) or None
            p_group = normalize_str(item.get("parallel_group") or item.get("group")) or None
            next_step = normalize_str(item.get("next_step") or item.get("next")) or None

            steps.append(StepRow(
                step_id=step_id,
                step=step_text,
                responsible=resp,
                type=t_val,
                if_yes=if_yes,
                if_no=if_no,
                parallel_group=p_group,
                next_step=next_step,
                description=normalize_str(item.get("description")) or None,
                system=normalize_str(item.get("system")) or None,
                input_data=normalize_str(item.get("input_data")) or None,
                output_data=normalize_str(item.get("output_data")) or None,
                duration=normalize_str(item.get("duration")) or None,
                row_number=idx + 1,
                sheet_name="JSON"
            ))

        return self.build_process_ir(steps, process_name=process_name, doc_id=filename)

    def validate_rows(self, steps: List[StepRow]) -> Tuple[bool, List[RowIssue], List[RowIssue]]:
        """
        Row-level validation returning plain-language messages with the row number.
        - Decision rows need both If Yes and If No targets.
        - Non-decision rows must leave them blank.
        - Every referenced Step ID must exist.
        - No self-references.
        - Unreachable rows are reported.
        - Parallel groups need at least two members and cannot contain Decision rows.
        - Blank Responsible offers an 'Unassigned' lane as a one-click fix.
        - Loops are allowed but listed for confirmation (INFO / WARNING).
        """
        errors: List[RowIssue] = []
        warnings: List[RowIssue] = []

        if not steps:
            errors.append(RowIssue(row=1, column="Step", message="Process table is empty. At least one step is required."))
            return False, errors, warnings

        # Index step IDs and locations
        step_id_map: Dict[str, StepRow] = {}
        for s in steps:
            s_id = s.step_id.strip()
            if s_id in step_id_map:
                errors.append(RowIssue(
                    row=s.row_number,
                    column="Step ID",
                    message=f"Row {s.row_number}: Duplicate Step ID '{s_id}' (first defined on row {step_id_map[s_id].row_number})."
                ))
            else:
                step_id_map[s_id] = s

        # Group parallel groups
        parallel_groups: Dict[str, List[StepRow]] = {}
        for s in steps:
            if s.parallel_group and s.parallel_group.strip():
                pg_key = s.parallel_group.strip()
                parallel_groups.setdefault(pg_key, []).append(s)

        ALLOWED_TYPES = {"task", "decision", "gateway", "exclusivegateway", "end", "endevent", "user", "service", "script", "send", "receive", "manual", "businessrule"}

        # Validate each row
        for s in steps:
            r = s.row_number
            s_id = s.step_id.strip()
            raw_type = s.type.strip().lower()
            is_decision = raw_type in ("decision", "gateway", "exclusivegateway")
            is_end = raw_type in ("end", "endevent")

            # 0. Valid type check
            if raw_type not in ALLOWED_TYPES:
                errors.append(RowIssue(
                    row=r,
                    column="Type",
                    message=f"Row {r}: Invalid type '{s.type}'. Must be Task, Decision, or End.",
                    severity="ERROR",
                    fix="Set Type to 'Task'"
                ))

            # 1. Blank Responsible check
            if not s.responsible or not s.responsible.strip():
                errors.append(RowIssue(
                    row=r,
                    column="Responsible",
                    message=f"Row {r}: Responsible is required for step '{s.step}'.",
                    severity="ERROR",
                    fix="Set Responsible to 'Unassigned'"
                ))
            elif s.responsible.strip().lower() in ("unassigned", "none"):
                warnings.append(RowIssue(
                    row=r,
                    column="Responsible",
                    message=f"Row {r}: Step '{s.step}' has no Responsible role assigned.",
                    severity="WARNING",
                    fix="Set Responsible to 'Unassigned'"
                ))

            # 2. Decision validation
            if is_decision:
                if not s.if_yes or not s.if_no:
                    errors.append(RowIssue(
                        row=r,
                        column="If Yes / If No",
                        message=f"Row {r}: Decision '{s.step}' requires both 'If Yes' and 'If No' target Step IDs."
                    ))
            else:
                if s.if_yes or s.if_no:
                    errors.append(RowIssue(
                        row=r,
                        column="If Yes / If No",
                        message=f"Row {r}: Step '{s.step}' is not a Decision; 'If Yes' and 'If No' must be blank."
                    ))

            # 3. Target references existence & self-reference checks
            targets_to_check: List[Tuple[str, str]] = []
            if s.if_yes:
                targets_to_check.append(("If Yes", s.if_yes.strip()))
            if s.if_no:
                targets_to_check.append(("If No", s.if_no.strip()))
            if s.next_step and s.next_step.strip().upper() != "END":
                targets_to_check.append(("Next Step", s.next_step.strip()))

            for col_name, tgt in targets_to_check:
                if tgt == s_id:
                    errors.append(RowIssue(
                        row=r,
                        column=col_name,
                        message=f"Row {r}: Step '{s.step}' points to itself; cannot reference itself as successor."
                    ))
                elif tgt not in step_id_map and tgt not in parallel_groups:
                    errors.append(RowIssue(
                        row=r,
                        column=col_name,
                        message=f"Row {r}: Referenced Step ID '{tgt}' does not exist in the process table."
                    ))

            # 4. Parallel group validation on row
            if s.parallel_group and s.parallel_group.strip():
                if is_decision:
                    errors.append(RowIssue(
                        row=r,
                        column="Parallel Group",
                        message=f"Row {r}: Decision steps cannot be placed inside a parallel group."
                    ))

            # 5. Values that would be silently ignored are reported instead
            if is_decision and s.next_step and s.next_step.strip():
                warnings.append(RowIssue(
                    row=r,
                    column="Next Step",
                    message=(
                        f"Row {r}: Decision '{s.step}' has a Next Step ('{s.next_step.strip()}') that is ignored; "
                        f"decisions branch only via 'If Yes' and 'If No'."
                    ),
                    severity="WARNING",
                    fix="Clear Next Step"
                ))
            if is_end and s.next_step and s.next_step.strip():
                warnings.append(RowIssue(
                    row=r,
                    column="Next Step",
                    message=f"Row {r}: End step '{s.step}' has a Next Step that is ignored; an End step terminates the branch.",
                    severity="WARNING",
                    fix="Clear Next Step"
                ))

        # 6. Parallel group member count + contiguity check
        row_index = {s.step_id.strip(): i for i, s in enumerate(steps)}
        for pg_name, members in parallel_groups.items():
            if len(members) < 2:
                for m in members:
                    errors.append(RowIssue(
                        row=m.row_number,
                        column="Parallel Group",
                        message=f"Row {m.row_number}: Parallel group '{pg_name}' must contain at least 2 steps (found {len(members)})."
                    ))
                continue
            positions = sorted(row_index[m.step_id.strip()] for m in members)
            if positions[-1] - positions[0] + 1 != len(positions):
                # A step between two members would fall through (by row order) into the middle
                # of the group and bypass the fork gateway, so the join could never complete.
                for m in members:
                    errors.append(RowIssue(
                        row=m.row_number,
                        column="Parallel Group",
                        message=(
                            f"Row {m.row_number}: Members of parallel group '{pg_name}' must be on consecutive rows; "
                            f"move the rows together or give the intervening steps an explicit Next Step."
                        ),
                        fix="Move group rows together"
                    ))

        # 6. Reachability and Loop analysis
        # Build adjacency graph for step IDs
        adj: Dict[str, List[str]] = {s.step_id.strip(): [] for s in steps}

        def add_target_to_adj(source_id: str, target_val: str):
            if target_val in parallel_groups:
                for mem in parallel_groups[target_val]:
                    adj[source_id].append(mem.step_id.strip())
            elif target_val in step_id_map:
                adj[source_id].append(target_val)
                target_step = step_id_map[target_val]
                if target_step.parallel_group and target_step.parallel_group.strip() in parallel_groups:
                    for mem in parallel_groups[target_step.parallel_group.strip()]:
                        adj[source_id].append(mem.step_id.strip())

        for idx, s in enumerate(steps):
            s_id = s.step_id.strip()
            is_decision = s.type.strip().lower() in ("decision", "gateway", "exclusivegateway")
            is_end = s.type.strip().lower() in ("end", "endevent")

            if is_end:
                continue

            if is_decision:
                if s.if_yes:
                    add_target_to_adj(s_id, s.if_yes.strip())
                if s.if_no:
                    add_target_to_adj(s_id, s.if_no.strip())
            elif s.next_step and s.next_step.strip().upper() == "END":
                continue
            elif s.next_step:
                add_target_to_adj(s_id, s.next_step.strip())
            else:
                # Default next row
                if idx + 1 < len(steps):
                    next_s = steps[idx + 1]
                    add_target_to_adj(s_id, next_s.step_id.strip())

        # If any member of a parallel group is reached, all members of the group are reached
        for pg_name, members in parallel_groups.items():
            for m1 in members:
                for m2 in members:
                    if m1.step_id.strip() != m2.step_id.strip():
                        adj[m1.step_id.strip()].append(m2.step_id.strip())

        # Reachability from step 1
        start_id = steps[0].step_id.strip()
        visited: Set[str] = set()
        queue = [start_id]
        while queue:
            curr = queue.pop(0)
            if curr not in visited:
                visited.add(curr)
                for neighbor in adj.get(curr, []):
                    if neighbor not in visited:
                        queue.append(neighbor)

        for s in steps:
            s_id = s.step_id.strip()
            if s_id not in visited:
                errors.append(RowIssue(
                    row=s.row_number,
                    column="Step ID",
                    message=f"Row {s.row_number}: Step '{s.step}' is unreachable from the process start."
                ))

        # Loop detection (confirmation / warning).
        # Only genuine back edges are reported: an edge whose target is still on the DFS stack
        # when the edge is explored. Parallel-group sibling links exist in `adj` purely for
        # reachability and are excluded here, otherwise every parallel group and every forward
        # edge that happens to sit on a cycle would be flagged as a loop.
        sibling_pairs: Set[Tuple[str, str]] = set()
        for members in parallel_groups.values():
            member_ids = [m.step_id.strip() for m in members]
            for a in member_ids:
                for b in member_ids:
                    if a != b:
                        sibling_pairs.add((a, b))

        flow_adj: Dict[str, List[str]] = {}
        for src_id, targets in adj.items():
            seen_targets: List[str] = []
            for tgt_id in targets:
                # skip sibling links, self links (a parallel member defaulting to the next
                # row, which is its own group) and duplicates
                if (src_id, tgt_id) in sibling_pairs or tgt_id == src_id or tgt_id in seen_targets:
                    continue
                seen_targets.append(tgt_id)
            flow_adj[src_id] = seen_targets

        back_edges: List[Tuple[str, str]] = []
        state: Dict[str, int] = {}  # absent = unvisited, 1 = on stack, 2 = done

        def dfs(node: str) -> None:
            state[node] = 1
            for nxt in flow_adj.get(node, []):
                nxt_state = state.get(nxt, 0)
                if nxt_state == 1:
                    if (node, nxt) not in back_edges:
                        back_edges.append((node, nxt))
                elif nxt_state == 0:
                    dfs(nxt)
            state[node] = 2

        if steps:
            dfs(steps[0].step_id.strip())
            for s_node in steps:  # cover nodes not reachable from the first step
                if state.get(s_node.step_id.strip(), 0) == 0:
                    dfs(s_node.step_id.strip())

        for src_id, tgt_id in back_edges:
            src_step = step_id_map.get(src_id)
            if src_step:
                warnings.append(RowIssue(
                    row=src_step.row_number,
                    column="Next Step",
                    message=f"Row {src_step.row_number}: Loop detected from Step '{src_id}' back to Step '{tgt_id}' (allowed; verify exit condition).",
                    severity="INFO"
                ))

        is_valid = len(errors) == 0
        return is_valid, errors, warnings

    def build_process_ir(self, steps: List[StepRow], process_name: str, doc_id: str) -> ProcessIR:
        """
        Converts validated step rows into a deterministic ProcessIR structure.
        """
        is_valid, errors, warnings = self.validate_rows(steps)
        if not is_valid:
            raise RowValidationError(errors + warnings)

        clean_proc_name = process_name.strip() or "Process"
        pool_id = sanitize_ncname(f"Pool_{clean_proc_name}", "Pool_1")
        proc_id = sanitize_ncname(f"Process_{clean_proc_name}", "Process_1")

        # 1. Lanes: each unique Responsible value in first-appearance order
        unique_roles: List[str] = []
        for s in steps:
            role = s.responsible.strip() or "Unassigned"
            if role not in unique_roles:
                unique_roles.append(role)

        if not unique_roles:
            unique_roles = ["General"]

        lanes: List[Lane] = []
        role_to_lane_id: Dict[str, str] = {}
        for idx, role in enumerate(unique_roles, start=1):
            lane_id = sanitize_ncname(f"Lane_{role}", f"Lane_{idx}")
            lanes.append(Lane(id=lane_id, name=role))
            role_to_lane_id[role] = lane_id

        pool = Pool(id=pool_id, name=clean_proc_name, lanes=lanes)

        elements: List[FlowNode] = []
        flows: List[SequenceFlow] = []

        # 2. Auto Start Event before step 1
        first_step = steps[0]
        first_role = first_step.responsible.strip() or "Unassigned"
        first_lane_id = role_to_lane_id.get(first_role, lanes[0].id)

        start_event_id = sanitize_ncname(f"StartEvent_{clean_proc_name}", "StartEvent_1")
        start_node = FlowNode(
            id=start_event_id,
            name="Start",
            type="startEvent",
            laneId=first_lane_id,
            sourceRefs=[_source_ref(doc_id, first_step.sheet_name, first_step.row_number, "Process Start")]
        )
        elements.append(start_node)

        # Index step rows and groups
        step_id_to_row = {s.step_id.strip(): s for s in steps}
        step_entry_nodes: Dict[str, str] = {}  # step_id -> node_id to connect into
        step_exit_nodes: Dict[str, str] = {}   # step_id -> node_id to connect out of

        # Pre-process parallel groups
        # We find distinct groups and their first member's position
        parallel_groups_map: Dict[str, List[StepRow]] = {}
        for s in steps:
            if s.parallel_group and s.parallel_group.strip():
                pg_key = s.parallel_group.strip()
                parallel_groups_map.setdefault(pg_key, []).append(s)

        group_split_ids: Dict[str, str] = {}
        group_join_ids: Dict[str, str] = {}

        for pg_key, members in parallel_groups_map.items():
            first_m = members[0]
            last_m = members[-1]
            split_lane_id = role_to_lane_id.get(first_m.responsible.strip() or "Unassigned", lanes[0].id)
            join_lane_id = role_to_lane_id.get(last_m.responsible.strip() or "Unassigned", lanes[0].id)

            split_id = sanitize_ncname(f"ParallelSplit_{pg_key}", f"ParallelSplit_{first_m.step_id}")
            join_id = sanitize_ncname(f"ParallelJoin_{pg_key}", f"ParallelJoin_{last_m.step_id}")

            group_split_ids[pg_key] = split_id
            group_join_ids[pg_key] = join_id

            split_node = FlowNode(
                id=split_id,
                name=f"Split {pg_key}",
                type="parallelGateway",
                laneId=split_lane_id,
                sourceRefs=[_source_ref(doc_id, first_m.sheet_name, first_m.row_number, f"Parallel Group {pg_key}")]
            )
            join_node = FlowNode(
                id=join_id,
                name=f"Join {pg_key}",
                type="parallelGateway",
                laneId=join_lane_id,
                sourceRefs=[_source_ref(doc_id, last_m.sheet_name, last_m.row_number, f"Parallel Group {pg_key}")]
            )
            elements.append(split_node)
            elements.append(join_node)

        # 3. Create FlowNodes for all steps
        handled_group_firsts: Set[str] = set()

        for idx, s in enumerate(steps):
            s_id = s.step_id.strip()
            role = s.responsible.strip() or "Unassigned"
            lane_id = role_to_lane_id.get(role, lanes[0].id)
            is_decision = s.type.strip().lower() in ("decision", "gateway", "exclusivegateway")
            is_end = s.type.strip().lower() in ("end", "endevent")

            source_ref = _source_ref(doc_id, s.sheet_name, s.row_number, s.step)

            if is_decision:
                gw_id = sanitize_ncname(f"Gateway_{s_id}", f"Gateway_{idx + 1}")
                node = FlowNode(
                    id=gw_id,
                    name=s.step,
                    type="exclusiveGateway",
                    laneId=lane_id,
                    documentation=s.description or "",
                    sourceRefs=[source_ref]
                )
                elements.append(node)
                step_entry_nodes[s_id] = gw_id
                step_exit_nodes[s_id] = gw_id

            elif is_end:
                end_id = sanitize_ncname(f"EndEvent_{s_id}", f"EndEvent_{idx + 1}")
                node = FlowNode(
                    id=end_id,
                    name=s.step or "End",
                    type="endEvent",
                    laneId=lane_id,
                    documentation=s.description or "",
                    sourceRefs=[source_ref]
                )
                elements.append(node)
                step_entry_nodes[s_id] = end_id
                step_exit_nodes[s_id] = end_id

            else:
                # Standard Task
                task_id = sanitize_ncname(f"Task_{s_id}", f"Task_{idx + 1}")
                node = FlowNode(
                    id=task_id,
                    name=s.step,
                    type="task",
                    laneId=lane_id,
                    documentation=s.description or "",
                    sourceRefs=[source_ref]
                )
                elements.append(node)

                # If member of parallel group
                if s.parallel_group and s.parallel_group.strip():
                    pg_key = s.parallel_group.strip()
                    split_id = group_split_ids[pg_key]
                    join_id = group_join_ids[pg_key]

                    # Sequence flow: Split -> Member Task
                    flows.append(SequenceFlow(
                        id=sanitize_ncname(f"Flow_split_{pg_key}_to_{s_id}", f"Flow_split_{s_id}"),
                        sourceId=split_id,
                        targetId=task_id
                    ))
                    # Sequence flow: Member Task -> Join
                    flows.append(SequenceFlow(
                        id=sanitize_ncname(f"Flow_{s_id}_to_join_{pg_key}", f"Flow_{s_id}_join"),
                        sourceId=task_id,
                        targetId=join_id
                    ))

                    # For the group's first member, its entry point is the split gateway
                    if pg_key not in handled_group_firsts:
                        handled_group_firsts.add(pg_key)
                        step_entry_nodes[s_id] = split_id
                    else:
                        step_entry_nodes[s_id] = task_id

                    # Exit point for group's last member is join gateway
                    if s == parallel_groups_map[pg_key][-1]:
                        step_exit_nodes[s_id] = join_id
                    else:
                        step_exit_nodes[s_id] = join_id
                else:
                    step_entry_nodes[s_id] = task_id
                    step_exit_nodes[s_id] = task_id

        # 4. Connect Start Event to Step 1
        first_step_entry = step_entry_nodes[first_step.step_id.strip()]
        flows.append(SequenceFlow(
            id=sanitize_ncname(f"Flow_start_to_{first_step.step_id.strip()}", "Flow_start_1"),
            sourceId=start_event_id,
            targetId=first_step_entry
        ))

        # 5. Connect Sequence Flows between steps
        end_event_counter = 1
        handled_join_outflow: Set[str] = set()

        for idx, s in enumerate(steps):
            s_id = s.step_id.strip()
            exit_node_id = step_exit_nodes[s_id]
            is_decision = s.type.strip().lower() in ("decision", "gateway", "exclusivegateway")
            is_end = s.type.strip().lower() in ("end", "endevent")

            if is_end:
                # End events have no outgoing sequence flow
                continue

            if is_decision:
                # If Yes -> target step
                yes_target_id = s.if_yes.strip()
                yes_target_node = step_entry_nodes[yes_target_id]
                flows.append(SequenceFlow(
                    id=sanitize_ncname(f"Flow_{s_id}_yes_to_{yes_target_id}", f"Flow_{s_id}_yes"),
                    name="Yes",
                    condition="Yes",
                    sourceId=exit_node_id,
                    targetId=yes_target_node
                ))

                # If No -> target step
                no_target_id = s.if_no.strip()
                no_target_node = step_entry_nodes[no_target_id]
                flows.append(SequenceFlow(
                    id=sanitize_ncname(f"Flow_{s_id}_no_to_{no_target_id}", f"Flow_{s_id}_no"),
                    name="No",
                    condition="No",
                    sourceId=exit_node_id,
                    targetId=no_target_node
                ))

            elif s.parallel_group and s.parallel_group.strip():
                pg_key = s.parallel_group.strip()
                members = parallel_groups_map[pg_key]
                # Outflow from parallel group is emitted once by the group's last member (or designated next_step)
                if s == members[-1] and pg_key not in handled_join_outflow:
                    handled_join_outflow.add(pg_key)
                    join_id = group_join_ids[pg_key]
                    next_step_val = s.next_step.strip() if s.next_step else None

                    if next_step_val and next_step_val.upper() == "END":
                        # Connect join to end event
                        end_node_id = sanitize_ncname(f"EndEvent_join_{pg_key}", f"EndEvent_pg_{end_event_counter}")
                        end_event_counter += 1
                        elements.append(FlowNode(
                            id=end_node_id,
                            name="End",
                            type="endEvent",
                            laneId=role_to_lane_id.get(s.responsible.strip() or "Unassigned", lanes[0].id),
                            sourceRefs=[_source_ref(doc_id, s.sheet_name, s.row_number, "Parallel Group End")]
                        ))
                        flows.append(SequenceFlow(
                            id=sanitize_ncname(f"Flow_join_{pg_key}_to_end", f"Flow_join_{pg_key}_end"),
                            sourceId=join_id,
                            targetId=end_node_id
                        ))
                    elif next_step_val and next_step_val in step_entry_nodes:
                        flows.append(SequenceFlow(
                            id=sanitize_ncname(f"Flow_join_{pg_key}_to_{next_step_val}", f"Flow_join_{pg_key}_{next_step_val}"),
                            sourceId=join_id,
                            targetId=step_entry_nodes[next_step_val]
                        ))
                    else:
                        # Fall through to next step after group if available
                        next_step_after_group = None
                        last_member_idx = steps.index(members[-1])
                        if last_member_idx + 1 < len(steps):
                            next_step_after_group = steps[last_member_idx + 1]

                        if next_step_after_group:
                            nxt_id = next_step_after_group.step_id.strip()
                            flows.append(SequenceFlow(
                                id=sanitize_ncname(f"Flow_join_{pg_key}_to_{nxt_id}", f"Flow_join_{pg_key}_{nxt_id}"),
                                sourceId=join_id,
                                targetId=step_entry_nodes[nxt_id]
                            ))
                        else:
                            # Last step in entire process -> create endEvent
                            end_node_id = sanitize_ncname(f"EndEvent_join_{pg_key}", f"EndEvent_pg_{end_event_counter}")
                            end_event_counter += 1
                            elements.append(FlowNode(
                                id=end_node_id,
                                name="End",
                                type="endEvent",
                                laneId=role_to_lane_id.get(s.responsible.strip() or "Unassigned", lanes[0].id),
                                sourceRefs=[_source_ref(doc_id, s.sheet_name, s.row_number, "Parallel Group End")]
                            ))
                            flows.append(SequenceFlow(
                                id=sanitize_ncname(f"Flow_join_{pg_key}_to_end", f"Flow_join_{pg_key}_end"),
                                sourceId=join_id,
                                targetId=end_node_id
                            ))

            else:
                # Normal Task sequence flow
                next_step_val = s.next_step.strip() if s.next_step else None

                if next_step_val and next_step_val.upper() == "END":
                    end_node_id = sanitize_ncname(f"EndEvent_{s_id}_term", f"EndEvent_{end_event_counter}")
                    end_event_counter += 1
                    elements.append(FlowNode(
                        id=end_node_id,
                        name="End",
                        type="endEvent",
                        laneId=lane_id,
                        sourceRefs=[_source_ref(doc_id, s.sheet_name, s.row_number, f"Step {s_id} End")]
                    ))
                    flows.append(SequenceFlow(
                        id=sanitize_ncname(f"Flow_{s_id}_to_{end_node_id}", f"Flow_{s_id}_end"),
                        sourceId=exit_node_id,
                        targetId=end_node_id
                    ))
                elif next_step_val and next_step_val in step_entry_nodes:
                    target_entry = step_entry_nodes[next_step_val]
                    flows.append(SequenceFlow(
                        id=sanitize_ncname(f"Flow_{s_id}_to_{next_step_val}", f"Flow_{s_id}_{next_step_val}"),
                        sourceId=exit_node_id,
                        targetId=target_entry
                    ))
                elif idx + 1 < len(steps):
                    # Next consecutive step
                    next_s = steps[idx + 1]
                    target_entry = step_entry_nodes[next_s.step_id.strip()]
                    flows.append(SequenceFlow(
                        id=sanitize_ncname(f"Flow_{s_id}_to_{next_s.step_id.strip()}", f"Flow_{s_id}_{next_s.step_id.strip()}"),
                        sourceId=exit_node_id,
                        targetId=target_entry
                    ))
                else:
                    # Last step in process -> auto End Event
                    end_node_id = sanitize_ncname(f"EndEvent_{s_id}_auto", f"EndEvent_{end_event_counter}")
                    end_event_counter += 1
                    elements.append(FlowNode(
                        id=end_node_id,
                        name="End",
                        type="endEvent",
                        laneId=lane_id,
                        sourceRefs=[_source_ref(doc_id, s.sheet_name, s.row_number, "Process End")]
                    ))
                    flows.append(SequenceFlow(
                        id=sanitize_ncname(f"Flow_{s_id}_to_{end_node_id}", f"Flow_{s_id}_end"),
                        sourceId=exit_node_id,
                        targetId=end_node_id
                    ))

        # Open questions for any warnings
        open_questions: List[OpenQuestion] = [
            OpenQuestion(topic="Template Notice", question=w.message, suggestedAssumption=w.fix or "")
            for idx, w in enumerate(warnings)
        ]

        return ProcessIR(
            id=proc_id,
            name=clean_proc_name,
            pools=[pool],
            elements=elements,
            flows=flows,
            openQuestions=open_questions
        )
